"""Admin API tests: authentication, authorisation, and every route."""

from __future__ import annotations

import io
from collections.abc import Iterator

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from talabflow import repository
from talabflow.api import create_app
from talabflow.config import Settings
from talabflow.db import session_scope
from talabflow.models import OrderStatus, StaffRole

from .conftest import ADMIN_PASSWORD, ADMIN_USERNAME, STAFF_PASSWORD, STAFF_USERNAME


@pytest.fixture
def client(settings: Settings, engine, staff_users) -> Iterator[TestClient]:
    app = create_app(settings)
    # The app builds its own engine against the same URL; the schema already exists.
    with TestClient(app) as opened:
        yield opened


@pytest.fixture
def admin_headers(client: TestClient) -> dict[str, str]:
    response = client.post(
        "/v1/auth/token", json={"username": ADMIN_USERNAME, "password": ADMIN_PASSWORD}
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.fixture
def staff_headers(client: TestClient) -> dict[str, str]:
    response = client.post(
        "/v1/auth/token", json={"username": STAFF_USERNAME, "password": STAFF_PASSWORD}
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.fixture
def order_reference(session_factory: sessionmaker[Session]) -> str:
    with session_scope(session_factory) as session:
        customer = repository.get_or_create_customer(
            session,
            channel="scripted",
            channel_user_id="1001",
            chat_id="1001",
            display_name="Amina",
        )
        order = repository.create_order(
            session,
            customer=customer,
            service_type="Repair",
            details="The washing machine will not drain",
            contact_phone="0555123456",
            address="12 Rue Didouche Mourad, Algiers",
        )
        return order.reference


# --------------------------------------------------------------------- health


def test_healthz_is_public(client: TestClient) -> None:
    response = client.get("/healthz")
    assert response.status_code == 200
    assert response.json()["transport"] == "scripted"


def test_readyz_checks_the_database(client: TestClient) -> None:
    response = client.get("/readyz")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_every_response_carries_a_request_id(client: TestClient) -> None:
    assert client.get("/healthz").headers["X-Request-ID"]


def test_a_supplied_request_id_is_echoed(client: TestClient) -> None:
    response = client.get("/healthz", headers={"X-Request-ID": "trace-me"})
    assert response.headers["X-Request-ID"] == "trace-me"


# --------------------------------------------------------------------- auth


def test_valid_credentials_return_a_token(client: TestClient) -> None:
    response = client.post(
        "/v1/auth/token", json={"username": ADMIN_USERNAME, "password": ADMIN_PASSWORD}
    )
    body = response.json()
    assert response.status_code == 200
    assert body["token_type"] == "bearer"
    assert body["role"] == "admin"
    assert body["expires_in_seconds"] > 0
    assert body["access_token"].count(".") == 2


def test_the_username_is_case_insensitive(client: TestClient) -> None:
    response = client.post(
        "/v1/auth/token", json={"username": ADMIN_USERNAME.upper(), "password": ADMIN_PASSWORD}
    )
    assert response.status_code == 200


@pytest.mark.parametrize(
    ("username", "password"),
    [
        (ADMIN_USERNAME, "the-wrong-password"),
        ("no-such-user", ADMIN_PASSWORD),
        (ADMIN_USERNAME, ""),
    ],
)
def test_bad_credentials_are_refused_with_one_identical_message(
    client: TestClient, username: str, password: str
) -> None:
    """Distinguishing "no such user" from "wrong password" lets an attacker enumerate valid
    usernames."""
    response = client.post("/v1/auth/token", json={"username": username, "password": password})
    assert response.status_code in (401, 422)
    if response.status_code == 401:
        assert response.json()["detail"] == "invalid username or password"


def test_protected_routes_reject_a_missing_token(client: TestClient) -> None:
    for method, path in [
        ("get", "/v1/orders"),
        ("get", "/v1/stats"),
        ("get", "/v1/staff"),
        ("get", "/v1/orders-export"),
    ]:
        response = getattr(client, method)(path)
        assert response.status_code == 401, path


def test_a_garbage_token_is_rejected(client: TestClient) -> None:
    response = client.get("/v1/orders", headers={"Authorization": "Bearer not-a-real-token"})
    assert response.status_code == 401


def test_a_token_signed_with_another_secret_is_rejected(
    client: TestClient, settings: Settings
) -> None:
    from talabflow.security import create_access_token

    forged = create_access_token(
        subject="attacker", role="admin", secret="a-completely-different-secret-value"
    )
    response = client.get("/v1/orders", headers={"Authorization": f"Bearer {forged}"})
    assert response.status_code == 401


# --------------------------------------------------------------------- rbac


def test_staff_can_read_orders(client: TestClient, staff_headers: dict[str, str]) -> None:
    assert client.get("/v1/orders", headers=staff_headers).status_code == 200


def test_staff_cannot_manage_staff_accounts(
    client: TestClient, staff_headers: dict[str, str]
) -> None:
    assert client.get("/v1/staff", headers=staff_headers).status_code == 403
    response = client.post(
        "/v1/staff",
        headers=staff_headers,
        json={"username": "sneaky", "password": "a-long-enough-password"},
    )
    assert response.status_code == 403


def test_an_admin_can_manage_staff_accounts(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    assert client.get("/v1/staff", headers=admin_headers).status_code == 200
    response = client.post(
        "/v1/staff",
        headers=admin_headers,
        json={"username": "new-tech", "password": "a-long-enough-password", "role": "staff"},
    )
    assert response.status_code == 201, response.text
    assert response.json()["username"] == "new-tech"
    assert response.json()["role"] == "staff"


def test_the_new_staff_user_can_log_in(client: TestClient, admin_headers: dict[str, str]) -> None:
    client.post(
        "/v1/staff",
        headers=admin_headers,
        json={"username": "new-tech", "password": "a-long-enough-password"},
    )
    response = client.post(
        "/v1/auth/token", json={"username": "new-tech", "password": "a-long-enough-password"}
    )
    assert response.status_code == 200
    assert response.json()["role"] == "staff"


def test_a_duplicate_username_is_a_conflict(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    payload = {"username": "dup-user", "password": "a-long-enough-password"}
    assert client.post("/v1/staff", headers=admin_headers, json=payload).status_code == 201
    assert client.post("/v1/staff", headers=admin_headers, json=payload).status_code == 409


def test_a_weak_password_is_refused(client: TestClient, admin_headers: dict[str, str]) -> None:
    response = client.post(
        "/v1/staff", headers=admin_headers, json={"username": "weak-user", "password": "short"}
    )
    assert response.status_code == 422


def test_an_invalid_username_shape_is_refused(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    for username in ["has spaces", "has/slash", "a", "e" * 100]:
        response = client.post(
            "/v1/staff",
            headers=admin_headers,
            json={"username": username, "password": "a-long-enough-password"},
        )
        assert response.status_code == 422, username


def test_the_staff_listing_never_includes_password_hashes(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    response = client.get("/v1/staff", headers=admin_headers)
    assert "scrypt" not in response.text
    assert all("password" not in key for row in response.json() for key in row)


# --------------------------------------------------------------------- orders


def test_list_orders_returns_a_page_with_a_total(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    body = client.get("/v1/orders", headers=staff_headers).json()
    assert body["total"] == 1
    assert body["items"][0]["reference"] == order_reference
    assert body["items"][0]["customer"]["display_name"] == "Amina"


def test_list_orders_filters_and_searches(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    assert client.get("/v1/orders?status=new", headers=staff_headers).json()["total"] == 1
    assert client.get("/v1/orders?status=completed", headers=staff_headers).json()["total"] == 0
    found = client.get("/v1/orders?search=washing", headers=staff_headers).json()
    assert found["total"] == 1


def test_list_orders_validates_its_query_parameters(
    client: TestClient, staff_headers: dict[str, str]
) -> None:
    assert client.get("/v1/orders?status=nonsense", headers=staff_headers).status_code == 422
    assert client.get("/v1/orders?limit=0", headers=staff_headers).status_code == 422
    assert client.get("/v1/orders?limit=9999", headers=staff_headers).status_code == 422
    assert client.get("/v1/orders?offset=-1", headers=staff_headers).status_code == 422


def test_get_order_includes_the_audit_trail_and_allowed_next_steps(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    body = client.get(f"/v1/orders/{order_reference}", headers=staff_headers).json()
    assert body["reference"] == order_reference
    assert len(body["events"]) == 1
    assert body["events"][0]["to_status"] == "new"
    assert body["events"][0]["actor"] == "customer"
    # Telling the client what is possible next means it never has to guess.
    assert sorted(body["allowed_next_statuses"]) == ["cancelled", "confirmed"]


def test_get_order_accepts_a_lower_case_reference(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    response = client.get(f"/v1/orders/{order_reference.lower()}", headers=staff_headers)
    assert response.status_code == 200


def test_an_unknown_reference_is_404(client: TestClient, staff_headers: dict[str, str]) -> None:
    assert client.get("/v1/orders/TF-20200101-AAAA", headers=staff_headers).status_code == 404


def test_a_valid_status_change_succeeds_and_records_the_actor(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    response = client.post(
        f"/v1/orders/{order_reference}/status",
        headers=staff_headers,
        json={"to_status": "confirmed", "note": "Technician assigned"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "confirmed"
    assert body["events"][-1]["actor"] == STAFF_USERNAME
    assert body["events"][-1]["note"] == "Technician assigned"


def test_an_invalid_status_change_is_a_conflict_not_a_bad_request(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    """The request is well-formed; it conflicts with the order's current state."""
    response = client.post(
        f"/v1/orders/{order_reference}/status",
        headers=staff_headers,
        json={"to_status": "completed"},
    )
    assert response.status_code == 409
    assert "from new to completed" in response.json()["detail"]


def test_an_unknown_status_value_is_refused(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    response = client.post(
        f"/v1/orders/{order_reference}/status",
        headers=staff_headers,
        json={"to_status": "teleported"},
    )
    assert response.status_code == 422


def test_an_over_long_note_is_refused(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    response = client.post(
        f"/v1/orders/{order_reference}/status",
        headers=staff_headers,
        json={"to_status": "confirmed", "note": "x" * 1000},
    )
    assert response.status_code == 422


def test_a_status_change_queues_a_customer_notification(
    client: TestClient,
    staff_headers: dict[str, str],
    order_reference: str,
    session_factory: sessionmaker[Session],
) -> None:
    client.post(
        f"/v1/orders/{order_reference}/status",
        headers=staff_headers,
        json={"to_status": "confirmed"},
    )
    with session_scope(session_factory) as session:
        due = repository.due_outbox_messages(session, limit=10)
        assert len(due) == 1
        assert order_reference in due[0].body


def test_stats_counts_every_status(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    body = client.get("/v1/stats", headers=staff_headers).json()
    assert body["total_orders"] == 1
    assert body["orders_by_status"]["new"] == 1
    assert body["orders_by_status"]["completed"] == 0
    assert set(body["orders_by_status"]) == {status.value for status in OrderStatus}


# --------------------------------------------------------------------- export


def test_xlsx_export_returns_a_real_workbook(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    response = client.get("/v1/orders-export?format=xlsx", headers=staff_headers)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert ".xlsx" in response.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.cell(row=2, column=1).value == order_reference


def test_csv_export_returns_a_bom_prefixed_file(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    response = client.get("/v1/orders-export?format=csv", headers=staff_headers)
    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")
    assert order_reference in response.content.decode("utf-8-sig")


def test_export_respects_a_status_filter(
    client: TestClient, staff_headers: dict[str, str], order_reference: str
) -> None:
    response = client.get("/v1/orders-export?format=csv&status=completed", headers=staff_headers)
    assert order_reference not in response.content.decode("utf-8-sig")


def test_an_unknown_export_format_is_refused(
    client: TestClient, staff_headers: dict[str, str]
) -> None:
    assert client.get("/v1/orders-export?format=pdf", headers=staff_headers).status_code == 422


def test_export_requires_authentication(client: TestClient) -> None:
    assert client.get("/v1/orders-export").status_code == 401


# --------------------------------------------------------------------- rate limiting


def test_the_api_rate_limit_returns_429(settings: Settings, engine, staff_users) -> None:
    tight = settings.model_copy(update={"api_rate_limit_per_minute": 3})
    with TestClient(create_app(tight)) as client:
        token = client.post(
            "/v1/auth/token", json={"username": STAFF_USERNAME, "password": STAFF_PASSWORD}
        ).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
        statuses = [client.get("/v1/orders", headers=headers).status_code for _ in range(6)]
        assert 429 in statuses
        limited = client.get("/v1/orders", headers=headers)
        assert limited.status_code == 429
        assert int(limited.headers["Retry-After"]) >= 1


def test_repeated_login_attempts_are_throttled(settings: Settings, engine, staff_users) -> None:
    """Regression guard: the login endpoint used to accept unlimited password guesses.

    The general rate limit only applies *after* a token has been verified, so it never protected
    the one endpoint that takes a password.
    """
    limited = settings.model_copy(update={"login_attempts_per_minute": 3})
    with TestClient(create_app(limited)) as client:
        codes = [
            client.post(
                "/v1/auth/token", json={"username": ADMIN_USERNAME, "password": f"wrong-guess-{n}"}
            ).status_code
            for n in range(5)
        ]
        assert codes == [401, 401, 401, 429, 429]

        # Once throttled, even the right password is refused -- otherwise the limit would only
        # slow an attacker down until the guess that mattered.
        blocked = client.post(
            "/v1/auth/token", json={"username": ADMIN_USERNAME, "password": ADMIN_PASSWORD}
        )
        assert blocked.status_code == 429
        assert int(blocked.headers["Retry-After"]) >= 1

        # The limit is per account: another member of staff can still sign in.
        other = client.post(
            "/v1/auth/token", json={"username": STAFF_USERNAME, "password": STAFF_PASSWORD}
        )
        assert other.status_code == 200


def test_the_login_throttle_ignores_username_case_and_padding(
    settings: Settings, engine, staff_users
) -> None:
    """``Admin-User`` and `` admin-user `` are the same account, so they share one budget."""
    limited = settings.model_copy(update={"login_attempts_per_minute": 2})
    with TestClient(create_app(limited)) as client:
        variants = [ADMIN_USERNAME, ADMIN_USERNAME.upper(), f" {ADMIN_USERNAME} "]
        codes = [
            client.post(
                "/v1/auth/token", json={"username": name, "password": "wrong-guess-x"}
            ).status_code
            for name in variants
        ]
        assert codes == [401, 401, 429]


@pytest.mark.parametrize("secret", ["change-me", "short-secret"])
def test_the_api_refuses_to_start_with_a_guessable_signing_secret(
    settings: Settings, engine, secret: str
) -> None:
    """Regression guard: the placeholder secret used to be accepted outside production.

    ``environment`` defaults to ``development``, so a deployment that never set it ran with the
    signing secret printed in ``.env.example`` -- and anyone who knew an admin's username could
    mint themselves a valid admin token without a password.
    """
    weak = settings.model_copy(update={"jwt_secret": secret, "environment": "development"})
    with pytest.raises(RuntimeError, match="TALABFLOW_JWT_SECRET"):
        create_app(weak)


def test_openapi_schema_is_generated(client: TestClient) -> None:
    paths = client.get("/openapi.json").json()["paths"]
    for expected in ["/v1/auth/token", "/v1/orders", "/v1/stats", "/v1/staff"]:
        assert expected in paths


# --------------------------------------------------------------- token revocation


def test_a_deactivated_account_loses_access_immediately(
    client: TestClient, admin_headers: dict[str, str], session_factory: sessionmaker[Session]
) -> None:
    """A valid signature is not enough on its own.

    Tokens last up to an hour. If the claims were trusted alone, deactivating a staff member would
    not actually revoke anything until their token happened to expire -- and revocation that takes
    effect "within an hour" is not revocation. The account is therefore re-checked per request.
    """
    assert client.get("/v1/orders", headers=admin_headers).status_code == 200

    with session_scope(session_factory) as session:
        user = repository.get_staff_user(session, ADMIN_USERNAME)
        assert user is not None
        user.is_active = False

    response = client.get("/v1/orders", headers=admin_headers)
    assert response.status_code == 401
    assert "no longer active" in response.json()["detail"]


def test_an_admin_can_deactivate_a_colleague_and_their_token_stops_working(
    client: TestClient, admin_headers: dict[str, str], staff_headers: dict[str, str]
) -> None:
    """The revocation the API promises has to be reachable through the API.

    Access was already re-checked on every request, but the only way to *trigger* a deactivation
    was to edit the database by hand.
    """
    assert client.get("/v1/orders", headers=staff_headers).status_code == 200

    response = client.patch(
        f"/v1/staff/{STAFF_USERNAME}", json={"is_active": False}, headers=admin_headers
    )
    assert response.status_code == 200, response.text
    assert response.json()["is_active"] is False

    assert client.get("/v1/orders", headers=staff_headers).status_code == 401
    login = client.post(
        "/v1/auth/token", json={"username": STAFF_USERNAME, "password": STAFF_PASSWORD}
    )
    assert login.status_code == 401


def test_an_admin_can_change_a_role_and_reset_a_password(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    response = client.patch(
        f"/v1/staff/{STAFF_USERNAME}",
        json={"role": "admin", "password": "a-brand-new-password"},
        headers=admin_headers,
    )
    assert response.status_code == 200
    assert response.json()["role"] == "admin"
    assert "password" not in response.text

    old = client.post(
        "/v1/auth/token", json={"username": STAFF_USERNAME, "password": STAFF_PASSWORD}
    )
    new = client.post(
        "/v1/auth/token", json={"username": STAFF_USERNAME, "password": "a-brand-new-password"}
    )
    assert (old.status_code, new.status_code) == (401, 200)
    assert new.json()["role"] == "admin"


def test_staff_cannot_update_accounts(client: TestClient, staff_headers: dict[str, str]) -> None:
    response = client.patch(
        f"/v1/staff/{ADMIN_USERNAME}", json={"is_active": False}, headers=staff_headers
    )
    assert response.status_code == 403


def test_the_only_admin_cannot_lock_everyone_out(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    for change in ({"is_active": False}, {"role": "staff"}):
        response = client.patch(f"/v1/staff/{ADMIN_USERNAME}", json=change, headers=admin_headers)
        assert response.status_code == 409
        assert "only active admin" in response.json()["detail"]
    assert client.get("/v1/staff", headers=admin_headers).status_code == 200


def test_updating_staff_validates_its_input(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    unknown = client.patch("/v1/staff/nobody", json={"is_active": False}, headers=admin_headers)
    empty = client.patch(f"/v1/staff/{STAFF_USERNAME}", json={}, headers=admin_headers)
    weak = client.patch(
        f"/v1/staff/{STAFF_USERNAME}", json={"password": "short"}, headers=admin_headers
    )
    assert (unknown.status_code, empty.status_code, weak.status_code) == (404, 422, 422)


def test_a_deleted_account_loses_access_immediately(
    client: TestClient, staff_headers: dict[str, str], session_factory: sessionmaker[Session]
) -> None:
    assert client.get("/v1/orders", headers=staff_headers).status_code == 200

    with session_scope(session_factory) as session:
        user = repository.get_staff_user(session, STAFF_USERNAME)
        assert user is not None
        session.delete(user)

    assert client.get("/v1/orders", headers=staff_headers).status_code == 401


def test_a_demotion_takes_effect_on_the_next_request(
    client: TestClient, admin_headers: dict[str, str], session_factory: sessionmaker[Session]
) -> None:
    """Authorisation reads the stored role, not the token's copy of it.

    Otherwise an admin demoted to staff would keep admin powers for the life of their token.
    """
    assert client.get("/v1/staff", headers=admin_headers).status_code == 200

    with session_scope(session_factory) as session:
        user = repository.get_staff_user(session, ADMIN_USERNAME)
        assert user is not None
        user.role = StaffRole.STAFF

    # The token still says "admin"; the database says otherwise, and the database wins.
    assert client.get("/v1/staff", headers=admin_headers).status_code == 403
    # ...but ordinary staff access still works.
    assert client.get("/v1/orders", headers=admin_headers).status_code == 200


def test_a_promotion_also_takes_effect_on_the_next_request(
    client: TestClient, staff_headers: dict[str, str], session_factory: sessionmaker[Session]
) -> None:
    assert client.get("/v1/staff", headers=staff_headers).status_code == 403

    with session_scope(session_factory) as session:
        user = repository.get_staff_user(session, STAFF_USERNAME)
        assert user is not None
        user.role = StaffRole.ADMIN

    assert client.get("/v1/staff", headers=staff_headers).status_code == 200
