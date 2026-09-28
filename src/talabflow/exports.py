"""CSV and XLSX export of orders.

The point of this module: "chat goes in, a spreadsheet comes out" is the outcome the buyer
actually asked for. Accounting runs on spreadsheets, not on an API.

Both writers stream into an in-memory buffer and share one column definition, so the two formats
can never drift apart.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from datetime import datetime
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .models import Order

#: ``(header, extractor)`` pairs. One definition, both formats.
COLUMNS: Final[list[tuple[str, str]]] = [
    ("Reference", "reference"),
    ("Status", "status"),
    ("Service", "service_type"),
    ("Details", "details"),
    ("Phone", "contact_phone"),
    ("Address", "address"),
    ("Customer", "customer_name"),
    ("Channel", "channel"),
    ("Created (UTC)", "created_at"),
    ("Updated (UTC)", "updated_at"),
]

_TIMESTAMP_FORMAT: Final = "%Y-%m-%d %H:%M:%S"


def _format_timestamp(value: datetime | None) -> str:
    return value.strftime(_TIMESTAMP_FORMAT) if value else ""


def order_to_row(order: Order) -> list[str]:
    """Flatten an order into export cells, in :data:`COLUMNS` order."""
    customer = order.customer
    values: dict[str, str] = {
        "reference": order.reference,
        "status": order.status.value,
        "service_type": order.service_type,
        "details": order.details,
        "contact_phone": order.contact_phone,
        "address": order.address,
        "customer_name": (customer.display_name if customer else None) or "",
        "channel": customer.channel if customer else "",
        "created_at": _format_timestamp(order.created_at),
        "updated_at": _format_timestamp(order.updated_at),
    }
    return [values[key] for _, key in COLUMNS]


def _sanitise_for_spreadsheet(value: str) -> str:
    """Neutralise formula injection in spreadsheet output.

    A customer can type ``=HYPERLINK("http://evil","click")`` or ``@SUM(A1)`` into a chat field.
    Excel and LibreOffice will happily evaluate that when staff open the export, which turns a
    free-text field into code execution on the buyer's machine (CSV/formula injection). Prefixing
    a leading formula trigger with an apostrophe forces the cell to stay text.
    """
    if value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def orders_to_csv(orders: Sequence[Order]) -> bytes:
    """Render orders as UTF-8 CSV with a BOM.

    The BOM is there so Excel on Windows opens Arabic text correctly instead of as mojibake --
    without it, the most common desktop spreadsheet mangles half the data in this application.
    """
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, quoting=csv.QUOTE_MINIMAL)
    writer.writerow([header for header, _ in COLUMNS])
    for order in orders:
        writer.writerow([_sanitise_for_spreadsheet(cell) for cell in order_to_row(order)])
    return buffer.getvalue().encode("utf-8-sig")


def orders_to_xlsx(orders: Sequence[Order], *, sheet_title: str = "Orders") -> bytes:
    """Render orders as a formatted XLSX workbook."""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # pragma: no cover - openpyxl always provides one
        sheet = workbook.create_sheet()
    sheet.title = sheet_title[:31]  # Excel's hard limit on sheet-name length

    headers = [header for header, _ in COLUMNS]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    for order in orders:
        sheet.append([_sanitise_for_spreadsheet(cell) for cell in order_to_row(order)])

    # Width from the widest value in each column, bounded so a long "details" field does not
    # produce a 500-character-wide column.
    for index, header in enumerate(headers, start=1):
        longest = max(
            [len(header)] + [len(str(sheet.cell(row=row, column=index).value or "")) for row in range(2, sheet.max_row + 1)]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 50)

    sheet.freeze_panes = "A2"
    if orders:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
